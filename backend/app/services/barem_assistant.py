"""AI assistant that drafts and edits a barem from an exam document + chat.

The model never writes the full barem schema. Getting `slot_id`s, criterion
ids, `part_label` wiring and `partial_credit_rule` syntax right is exactly
what makes hand-authoring hard, and an LLM asked to emit that schema directly
gets it subtly wrong in ways that only show up as a mis-graded class. Instead
it answers in a small *draft* format — questions, parts, criteria with a
type, a score, answers and a plain-language partial-credit rule — and the
frontend builds real `RubricQuestion`s from it with the same factory the
editor uses (`frontend/src/modules/barem/assistant/fromDraft.ts`).

Nothing here touches the stored barem: `/chat` returns a proposal, and the
teacher accepts it question by question in the browser.

Out of scope for now, and therefore *locked*: questions with printed tables
or `conditional_outputs` (answers that depend on the student's index). The
draft format can't express them, so rewriting one would silently drop that
structure — the model is told not to touch them and they're filtered out of
its answer regardless.
"""

from __future__ import annotations

import io
import json
import os
import re
from typing import Any

import requests
from dotenv import load_dotenv

# Same .env pipeline.py reads (repo root, found by searching upward). The API's
# own Settings only reads backend/.env, so outside Docker — where compose isn't
# there to inject the file — the LLM credentials would otherwise be missing.
load_dotenv()

MAX_DOCUMENT_CHARS = 60_000
MAX_HISTORY_MESSAGES = 12
DRAFT_TYPES = {"matching", "logical", "visual"}


class AssistantError(Exception):
    """A failure worth showing the teacher verbatim."""


# ── documents ───────────────────────────────────────────────────────────────


def extract_document_text(filename: str, data: bytes) -> str:
    """Plain text of a .pdf or .docx exam. Scanned PDFs are rejected, not guessed at."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:  # noqa: BLE001 - any parse failure is a bad upload
            raise AssistantError(f"Không đọc được file PDF: {exc}") from exc
        text = "\n\n".join(f"--- Trang {i} ---\n{p.strip()}" for i, p in enumerate(pages, 1) if p.strip())
        if len(text.replace("--- Trang", "").strip()) < 30:
            raise AssistantError(
                "PDF này không có chữ (có vẻ là ảnh scan). Hãy dùng file Word, PDF xuất từ Word, "
                "hoặc gõ đề vào khung chat."
            )
    elif name.endswith(".docx"):
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise AssistantError(f"Không đọc được file Word: {exc}") from exc
        blocks: list[str] = [p.text for p in document.paragraphs]
        # Tables hold half of a typical exam (printed I/O grids, answer keys),
        # and python-docx keeps them out of `paragraphs`.
        for index, table in enumerate(document.tables, 1):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            blocks.append(f"[Bảng {index}]\n" + "\n".join(rows))
        text = "\n".join(blocks).strip()
    else:
        raise AssistantError("Chỉ hỗ trợ file .pdf hoặc .docx (file .doc cũ hãy lưu lại thành .docx).")

    if not text:
        raise AssistantError("File không có nội dung chữ nào.")
    if len(text) > MAX_DOCUMENT_CHARS:
        text = text[:MAX_DOCUMENT_CHARS] + "\n[... đã cắt bớt phần cuối vì đề quá dài ...]"
    return text


# ── barem <-> draft ─────────────────────────────────────────────────────────


def _walk(criteria: list[dict]) -> list[dict]:
    out: list[dict] = []
    for criterion in criteria or []:
        out.append(criterion)
        out.extend(_walk(criterion.get("sub_criteria") or []))
    return out


def locked_questions(exam: dict) -> set[int]:
    locked: set[int] = set()
    for question in exam.get("teacher_barem") or []:
        parts = (question.get("question") or {}).get("parts") or []
        has_table = any(part.get("tables") for part in parts)
        has_conditional = any(c.get("conditional_outputs") for c in _walk(question.get("grading_rule") or []))
        if has_table or has_conditional:
            locked.add(int(question.get("question_number", 0)))
    return locked


def _partial_of(rule) -> list[dict]:
    rules = rule if isinstance(rule, list) else [rule] if rule else []
    out = []
    for r in rules:
        match = re.search(r">=\s*(\d+)", str(r.get("condition", "")))
        if r.get("type") == "count_correct_tokens" and match:
            out.append({"min_correct": int(match.group(1)), "score": r.get("partial_score")})
        else:
            out.append({"raw_rule": r})
    return out


def exam_to_draft(exam: dict) -> dict:
    """The current barem in the draft format, so the model edits what it can read.

    Lossy on purpose — wrapper criteria are flattened to their gradable leaves —
    which is why a rewritten question replaces the old one wholesale and why
    locked questions are shown only by number.
    """
    locked = locked_questions(exam)
    questions = []
    for q in exam.get("teacher_barem") or []:
        number = int(q.get("question_number", 0))
        if number in locked:
            questions.append({"number": number, "locked": True})
            continue
        parts = (q.get("question") or {}).get("parts") or []
        criteria = []
        for c in _walk(q.get("grading_rule") or []):
            if c.get("sub_criteria"):
                continue
            value = c.get("expected_value") if isinstance(c.get("expected_value"), dict) else {}
            item = {
                "part": c.get("part_label"),
                "type": c.get("question_type"),
                "score": c.get("score"),
                "content": c.get("content"),
            }
            optional = {
                "grader_note": c.get("grader_note"),
                "answers": c.get("expected_outputs"),
                "tokens": c.get("expected_output_tokens"),
                "partial": _partial_of(c.get("partial_credit_rule")) or None,
                "keywords": value.get("keywords"),
                "sample_solution": value.get("sample_solution"),
            }
            item.update({k: v for k, v in optional.items() if v})
            criteria.append(item)
        questions.append(
            {
                "number": number,
                "text": (q.get("question") or {}).get("text", ""),
                "parts": [{"label": p.get("part_label"), "text": p.get("text", "")} for p in parts],
                "criteria": criteria,
            }
        )
    return {
        "meta": {k: exam.get(k) for k in ("ma_de", "subject", "total_score")},
        "questions": questions,
    }


# ── prompt ──────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = """Bạn là trợ lý giúp giáo viên Việt Nam soạn BAREM CHẤM ĐIỂM cho một hệ thống chấm bài tự động bằng AI.
Giáo viên đưa đề thi (chữ trích từ file Word/PDF) và nói bằng lời đáp án, cách cho điểm. Bạn dựng barem ở ĐỊNH DẠNG NHÁP dưới đây.

## Mỗi câu hỏi
- "number": số thứ tự câu (số nguyên, đúng như trên đề).
- "text": đề bài NGUYÊN VĂN, kể cả đoạn code, kể cả phần in sẵn (chữ ký hàm, biến khai sẵn). Không tóm tắt.
- "parts": các ý của câu. Câu không chia ý: [{"label": "main", "text": ""}]. Câu có ý a, b, c: [{"label": "a", "text": "yêu cầu riêng của ý a"}, ...]. "text" của ý CHỈ ghi phần chưa có trong đề bài, không chép lại đề.
- "criteria": các tiêu chí chấm. Mỗi ý có ít nhất 1 tiêu chí; "part" phải trùng "label" của một ý.

## Mỗi tiêu chí
- "type" — CHỈ một trong ba:
  - "matching": đáp án là một giá trị/chuỗi CỤ THỂ, xác định (kết quả in ra màn hình, một con số, một từ). Bắt buộc có "answers" (danh sách các cách viết đúng TRỌN VẸN, mỗi phần tử là một đáp án đầy đủ). Nếu sai một phần vẫn được điểm: thêm "tokens" (các mảnh nhỏ theo thứ tự, ví dụ đáp án "3529" → ["3","5","2","9"]) và "partial": [{"min_correct": số mảnh đúng tối thiểu, "score": điểm}].
  - "logical": viết code, giải thích, chứng minh, câu tự luận — chấm theo logic, chấp nhận nhiều cách làm. Nên có "sample_solution" (lời giải mẫu) và "keywords" (CHỈ những thứ học sinh BẮT BUỘC phải tự viết, KHÔNG liệt kê tên hàm/biến đề đã in sẵn).
  - "visual": học sinh VẼ (lưu đồ, sơ đồ, hình). Chỉ cần "content" và "grader_note".
- "score": điểm tối đa của tiêu chí (số). Tổng điểm các tiêu chí của một câu = điểm của câu ghi trên đề.
- "content": mô tả tiêu chí cụ thể — nêu đáp án đúng và điều gì làm mất điểm.
- "grader_note" (tuỳ chọn): ngoại lệ, trường hợp đặc biệt, điều giáo viên dặn.

## Quy tắc
- KHÔNG bịa đáp án. Nếu đề/giáo viên không cho đáp án hoặc cách cho điểm của một câu, vẫn dựng câu đó với hiểu biết tốt nhất của bạn và GHI RÕ vào "confidence_notes" để giáo viên kiểm tra.
- Câu có BẢNG cần điền hoặc đáp án PHỤ THUỘC STT học sinh: KHÔNG dựng — ghi vào "confidence_notes" rằng câu đó giáo viên cần soạn tay.
- Các câu "locked": true trong barem hiện tại TUYỆT ĐỐI KHÔNG đưa vào "questions" hay "remove_questions".
- Khi giáo viên chỉ yêu cầu sửa vài câu, CHỈ trả lại những câu đó (mỗi câu trả về đầy đủ, nó sẽ thay thế câu cũ). Không trả lại câu không đổi.
- "remove_questions": số các câu giáo viên yêu cầu xoá.
- "meta": chỉ trả khi biết mã đề / môn / tổng điểm từ đề hoặc giáo viên.

## Định dạng trả lời
CHỈ MỘT object JSON hợp lệ, không markdown, không chữ nào ngoài JSON:
{"reply": "giải thích ngắn bằng tiếng Việt cho giáo viên",
 "meta": {"ma_de": "...", "subject": "...", "total_score": 10},
 "questions": [{"number": 1, "text": "...", "parts": [{"label": "main", "text": ""}],
   "criteria": [{"part": "main", "type": "matching", "score": 0.5, "content": "...",
     "answers": ["3529"], "tokens": ["3","5","2","9"], "partial": [{"min_correct": 2, "score": 0.25}]}]}],
 "remove_questions": [],
 "confidence_notes": ["..."]}
"""


def _context_block(exam: dict, document_text: str | None) -> str:
    draft = exam_to_draft(exam)
    parts = [
        "## Barem hiện tại (định dạng nháp; câu \"locked\" không được sửa)",
        json.dumps(draft, ensure_ascii=False),
    ]
    if document_text:
        parts += ["## Nội dung file đề thi giáo viên đã tải lên", document_text]
    else:
        parts.append("## Giáo viên chưa tải file đề — chỉ dựa vào lời giáo viên trong hội thoại.")
    return "\n\n".join(parts)


# ── LLM ─────────────────────────────────────────────────────────────────────


def _parse_json(text: str) -> dict | None:
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip())
    start = text.find("{")
    if start < 0:
        return None
    try:
        value, _ = json.JSONDecoder().raw_decode(text, start)
    except json.JSONDecodeError:
        return None
    return value if isinstance(value, dict) else None


def _call_llm(messages: list[dict]) -> str:
    api = os.environ.get("LLM_MODEL_API", "")
    model = os.environ.get("LLM_MODEL_NAME", "")
    key = os.environ.get("LLM_API_KEY", "")
    if not api or not model:
        raise AssistantError("Chưa cấu hình LLM (thiếu LLM_MODEL_API / LLM_MODEL_NAME trong .env).")
    headers = {"Content-Type": "application/json"}
    if key:
        headers["Authorization"] = f"Bearer {key}"
    try:
        response = requests.post(
            api,
            headers=headers,
            json={"model": model, "messages": messages, "temperature": 0.2, "max_tokens": 8000},
            timeout=300,
        )
    except requests.RequestException as exc:
        raise AssistantError(f"Không gọi được LLM: {exc}") from exc
    if response.status_code != 200:
        raise AssistantError(f"LLM trả lỗi HTTP {response.status_code}: {response.text[:300]}")
    try:
        return response.json()["choices"][0]["message"]["content"] or ""
    except (KeyError, IndexError, ValueError) as exc:
        raise AssistantError("LLM trả về dữ liệu không đúng định dạng.") from exc


def _normalise(raw: dict, locked: set[int]) -> dict:
    """Keep only what the frontend can build, and say what was dropped."""
    notes = [str(n) for n in raw.get("confidence_notes") or [] if n]
    questions = []
    for q in raw.get("questions") or []:
        if not isinstance(q, dict):
            continue
        try:
            number = int(q.get("number"))
        except (TypeError, ValueError):
            continue
        if number in locked:
            notes.append(f"Câu {number} có bảng / đáp án theo STT nên trợ lý không sửa — giữ nguyên bản cũ.")
            continue
        criteria = []
        for c in q.get("criteria") or []:
            if not isinstance(c, dict):
                continue
            if c.get("type") not in DRAFT_TYPES:
                notes.append(f"Câu {number}: trợ lý đề xuất kiểu chấm '{c.get('type')}' không hỗ trợ — đã đổi thành 'logical'.")
                c = {**c, "type": "logical"}
            criteria.append(c)
        questions.append({**q, "number": number, "criteria": criteria})
    removals = []
    for value in raw.get("remove_questions") or []:
        try:
            number = int(value)
        except (TypeError, ValueError):
            continue
        if number in locked:
            notes.append(f"Câu {number} bị khoá nên không xoá.")
        else:
            removals.append(number)
    meta = raw.get("meta") if isinstance(raw.get("meta"), dict) else None
    return {
        "reply": str(raw.get("reply") or ""),
        "meta": meta,
        "questions": questions,
        "remove_questions": removals,
        "confidence_notes": notes,
    }


def chat(messages: list[dict], exam: dict, document_text: str | None) -> dict:
    history = [
        {"role": m["role"], "content": str(m.get("content", ""))}
        for m in messages[-MAX_HISTORY_MESSAGES:]
        if m.get("role") in ("user", "assistant")
    ]
    if not history or history[-1]["role"] != "user":
        raise AssistantError("Tin nhắn cuối phải là của giáo viên.")

    llm_messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "system", "content": _context_block(exam, document_text)},
        *history,
    ]
    text = _call_llm(llm_messages)
    parsed = _parse_json(text)
    if parsed is None:
        # One retry: models occasionally wrap the JSON in prose despite the
        # instruction. Show it its own answer and ask for the JSON alone.
        retry = llm_messages + [
            {"role": "assistant", "content": text[:4000]},
            {"role": "user", "content": "Trả lời lại CHỈ bằng một object JSON hợp lệ đúng định dạng đã nêu, không có chữ nào khác."},
        ]
        parsed = _parse_json(_call_llm(retry))
    if parsed is None:
        raise AssistantError("Trợ lý trả lời không đúng định dạng. Hãy thử lại hoặc diễn đạt yêu cầu ngắn gọn hơn.")
    return _normalise(parsed, locked_questions(exam))
